"""Core logic for the Question Paper Generator.

Pipeline:
  read syllabus -> detect units (LLM) -> plan unit allocation (code)
  -> generate each set (LLM) -> validate rules (code, with retries)
  -> fill Word template (code) -> zip for download
"""
import copy
import difflib
import io
import json
import random
import re
import zipfile

import docx
from docx.oxml.ns import qn

# --------------------------------------------------------------------------
# Constants
# --------------------------------------------------------------------------
COMMAND_WORDS = [
    "Explain", "Discuss", "Describe", "Elaborate on", "List out", "Summarise",
    "Recall", "Analyse", "Illustrate", "Evaluate", "Examine", "Outline",
    "Compare", "Distinguish", "Differentiate", "Enumerate", "Trace", "Define",
    "Highlight", "State", "Justify", "Critically examine", "Suggest",
]

# regexes used to count command words in short + long questions
VERB_PATTERNS = {
    "explain": r"\bexplain\b", "discuss": r"\bdiscuss\b", "describe": r"\bdescribe\b",
    "elaborate on": r"\belaborate on\b", "list out": r"\blist out\b",
    "summarise": r"\bsummari[sz]e\b", "recall": r"\brecall\b",
    "analyse": r"\banaly[sz]e\b", "illustrate": r"\billustrate\b",
    "evaluate": r"\bevaluate\b", "examine": r"\bexamine\b", "outline": r"\boutline\b",
    "compare": r"\bcompare\b", "distinguish": r"\bdistinguish\b",
    "differentiate": r"\bdifferentiate\b", "enumerate": r"\benumerate\b",
    "trace": r"\btrace\b", "define": r"\bdefine\b", "highlight": r"\bhighlight\b",
    "state": r"\bstate\b", "justify": r"\bjustify\b", "suggest": r"\bsuggest\b",
    "what is/are": r"\bwhat (is|are)\b", "comment on": r"\bcomment on\b",
}

# options that point to other options are not allowed
BANNED_OPTION = re.compile(
    r"\b(all|none|any|both|either|neither)\b.{0,14}\b(above|these|options?|choices?)\b"
    r"|\b[1-4]\s*(and|&|,|or)\s*[1-4]\b"
    r"|\b[A-D]\s*(and|&)\s*[A-D]\b"
    r"|\(\s*[a-dA-D]\s*\)\s*(and|&)",
    re.I,
)

SYSTEM_PROMPT = (
    "You are an experienced university examiner who sets question papers for MBA courses "
    "offered by a distance and online education centre. You write accurate, unambiguous "
    "questions strictly from the syllabus provided. Respond with ONE JSON object only: "
    "no markdown fences and no commentary."
)

# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------

def read_syllabus(file_bytes: bytes, filename: str) -> str:
    name = filename.lower()
    if name.endswith(".docx"):
        d = docx.Document(io.BytesIO(file_bytes))
        parts = []
        body = d.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                t = "".join(x.text or "" for x in child.iter(qn("w:t")))
                if t.strip():
                    parts.append(t.strip())
            elif child.tag == qn("w:tbl"):
                for tr in child.iter(qn("w:tr")):
                    cells = []
                    for tc in tr.iter(qn("w:tc")):
                        cells.append(" ".join("".join(x.text or "" for x in p.iter(qn("w:t")))
                                              for p in tc.iter(qn("w:p"))).strip())
                    row = " | ".join(c for c in cells if c)
                    if row:
                        parts.append(row)
        return "\n".join(parts)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(file_bytes))
        return "\n".join((p.extract_text() or "") for p in r.pages)
    return file_bytes.decode("utf-8", errors="ignore")


# --------------------------------------------------------------------------
# LLM access (Gemini free tier or Claude API)
# --------------------------------------------------------------------------

def call_llm(provider: str, api_key: str, model: str, system: str, user: str,
             temperature: float = 0.7) -> str:
    if provider.startswith("Google Gemini"):
        from google import genai
        from google.genai import types
        client = genai.Client(api_key=api_key)
        resp = client.models.generate_content(
            model=model,
            contents=user,
            config=types.GenerateContentConfig(
                system_instruction=system,
                response_mime_type="application/json",
                temperature=temperature,
            ),
        )
        return resp.text
    if provider.startswith("Anthropic"):
        import anthropic
        client = anthropic.Anthropic(api_key=api_key)
        resp = client.messages.create(
            model=model, max_tokens=8000, temperature=temperature,
            system=system, messages=[{"role": "user", "content": user}],
        )
        return "".join(b.text for b in resp.content if getattr(b, "type", "") == "text")
    raise ValueError("Unknown provider")


def parse_json(text: str) -> dict:
    text = text.strip()
    text = re.sub(r"^```(?:json)?|```$", "", text, flags=re.M).strip()
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("No JSON object found in model output")
    return json.loads(text[start:end + 1])


# --------------------------------------------------------------------------
# Step 1: detect units
# --------------------------------------------------------------------------

def detect_units(llm, syllabus_text: str) -> dict:
    prompt = f"""Read this syllabus and list its units.

Return JSON exactly like:
{{"course_name": "...", "units": [{{"no": "I", "title": "...", "topics": "comma separated topics of the unit"}}]}}

Rules:
- Use Roman numerals (I, II, III ...) for "no", converting if the syllabus uses digits.
- Keep the unit order of the syllabus. Copy topics faithfully; do not invent topics.

SYLLABUS:
{syllabus_text[:60000]}"""
    return parse_json(llm(SYSTEM_PROMPT, prompt))


# --------------------------------------------------------------------------
# Step 2: plan which unit each question comes from (done in code, not by AI)
# --------------------------------------------------------------------------

def _pick(units, k, usage, rng, exclude=()):
    pool = [u for u in units if u not in exclude]
    rng.shuffle(pool)
    pool.sort(key=lambda u: usage[u])
    if len(pool) >= k:
        chosen = pool[:k]
    else:  # fewer units than questions: repeat units as little as possible
        chosen = [pool[i % len(pool)] for i in range(k)]
    for c in chosen:
        usage[c] += 1
    return chosen


def plan_sets(unit_ids, n_sets, seed=None):
    rng = random.Random(seed)
    mcq_use = {u: 0 for u in unit_ids}
    sl_use = {u: 0 for u in unit_ids}
    long_use = {u: 0 for u in unit_ids}
    plans = []
    for i in range(n_sets):
        mcq_units = _pick(unit_ids, 10, mcq_use, rng)
        sl_units = _pick(unit_ids, 10, sl_use, rng)
        long_units = _pick(sl_units, 4, long_use, rng)
        short_units = [u for u in sl_units if u not in long_units]
        if len(short_units) != 6:  # only when there are fewer than 10 units
            short_units = (short_units + sl_units)[:6]
        long_marks = ["5 + 5", "5 + 5", "10", "10"]
        rng.shuffle(long_marks)
        plans.append(dict(
            set_no=i + 1,
            mcq=sorted(mcq_units, key=unit_ids.index),
            short=sorted(short_units, key=unit_ids.index),
            long=sorted(long_units, key=unit_ids.index),
            long_marks=long_marks,
        ))
    return plans


# --------------------------------------------------------------------------
# Step 3: prompt for one set
# --------------------------------------------------------------------------

def build_prompt(course, units, plan, prior_questions, short_words, long_words, errors=None):
    used = sorted(set(plan["mcq"] + plan["short"] + plan["long"]), key=lambda x: [u["no"] for u in units].index(x))
    unit_text = "\n".join(
        f"Unit {u['no']}: {u['title']} -- {u['topics']}" for u in units if u["no"] in used
    )
    long_spec = "\n".join(
        f"  {i + 1}. Unit {u}: marks \"{m}\"" +
        (" (two parts, written as \"(a) ... (5 marks)\\n(b) ... (5 marks)\")" if m == "5 + 5"
         else " (a single 10-mark question)")
        for i, (u, m) in enumerate(zip(plan["long"], plan["long_marks"]))
    )
    prior = "\n".join(f"- {q}" for q in prior_questions[-120:]) or "(none yet)"
    fix = ""
    if errors:
        fix = "\nYOUR PREVIOUS ATTEMPT FAILED THESE CHECKS. FIX THEM ALL:\n" + "\n".join(f"- {e}" for e in errors) + "\n"
    return f"""Set question paper Set {plan['set_no']} for: {course}.

SYLLABUS UNITS TO USE:
{unit_text}

STRUCTURE (follow exactly, in this order):
- 10 MCQs, 2 marks each, from these units in order: {plan['mcq']}
- 6 short-answer questions, 5 marks each, answerable in about {short_words} words, from these units in order: {plan['short']}
- 4 long-answer questions, 10 marks each, answerable in about {long_words} words, from these units in order:
{long_spec}

RULES:
1. Every question must come only from the topics of its unit as listed above.
2. MCQs: exactly 4 options and exactly one correct answer. Put the index (1-4) of the correct option in "correct". Distractors must be plausible but clearly wrong.
   NEVER use options such as "All of the above", "None of the above", "Both A and B", "1 and 2" or any option that refers to other options. Every option must stand on its own.
3. Do not begin MCQ stems with command words like Explain or Discuss; use natural question stems.
4. Short and long questions must use VARIED command words from: {', '.join(COMMAND_WORDS)}.
   Each command word may appear at most TWICE in the whole paper (counting every part (a)/(b) of long questions); most should be used once.
5. Do not repeat or closely paraphrase any question already used in earlier sets:
{prior}
6. Language: clear, formal, suitable for MBA learners. No trick questions.
{fix}
Return JSON exactly like:
{{"mcq": [{{"unit": "I", "question": "...", "options": ["...", "...", "...", "..."], "correct": 2}}],
 "short": [{{"unit": "II", "question": "..."}}],
 "long": [{{"unit": "III", "question": "...", "marks": "10"}}]}}"""


# --------------------------------------------------------------------------
# Step 4: validate (code decides, not AI)
# --------------------------------------------------------------------------

def _norm(s):
    return re.sub(r"[^a-z0-9 ]", "", s.lower())


def validate_paper(paper, plan, prior_questions, unit_ids, max_verb_repeat=2):
    errs = []
    try:
        mcq, short, long_ = paper["mcq"], paper["short"], paper["long"]
    except KeyError as e:
        return [f"Missing key {e}"]
    if (len(mcq), len(short), len(long_)) != (10, 6, 4):
        return [f"Wrong counts: mcq={len(mcq)}, short={len(short)}, long={len(long_)} (need 10, 6, 4)"]

    for name, items, planned in (("mcq", mcq, plan["mcq"]), ("short", short, plan["short"]), ("long", long_, plan["long"])):
        got = [str(x.get("unit", "")).strip().upper() for x in items]
        if sorted(got) != sorted(planned):
            errs.append(f"{name} units {got} do not match the planned units {planned}")

    for i, q in enumerate(mcq, 1):
        opts = q.get("options", [])
        if len(opts) != 4 or len({_norm(o) for o in opts}) != 4:
            errs.append(f"MCQ {i}: needs 4 distinct options")
            continue
        if q.get("correct") not in (1, 2, 3, 4):
            errs.append(f"MCQ {i}: 'correct' must be 1-4")
        for o in opts:
            if BANNED_OPTION.search(o):
                errs.append(f"MCQ {i}: banned option style -> '{o}'")

    for i, (q, m) in enumerate(zip(long_, plan["long_marks"]), 1):
        text = q.get("question", "")
        if str(q.get("marks", "")).replace(" ", "") != m.replace(" ", ""):
            errs.append(f"Long {i}: marks should be '{m}'")
        if m == "5 + 5" and not ("(a)" in text and "(b)" in text):
            errs.append(f"Long {i}: must have parts (a) and (b)")

    sl_text = " ".join(q["question"] for q in short + long_).lower()
    for verb, pat in VERB_PATTERNS.items():
        n = len(re.findall(pat, sl_text))
        if n > max_verb_repeat:
            errs.append(f"Command word '{verb}' used {n} times (max {max_verb_repeat})")

    mine = [q["question"] for q in mcq + short + long_]
    for q in mine:
        for p in prior_questions:
            if difflib.SequenceMatcher(None, _norm(q), _norm(p)).ratio() > 0.85:
                errs.append(f"Too similar to an earlier question: '{q[:70]}'")
                break
    if len({_norm(q) for q in mine}) != len(mine):
        errs.append("Duplicate questions inside the paper")
    return errs


def finalise(paper, rng=None):
    """Shuffle MCQ options so the correct answer is spread across positions."""
    rng = rng or random.Random()
    out = []
    for q in paper["mcq"]:
        opts = list(q["options"])
        correct_text = opts[q["correct"] - 1]
        rng.shuffle(opts)
        out.append(dict(unit=str(q["unit"]).upper(), question=q["question"].strip(),
                        options=opts, correct=opts.index(correct_text) + 1))
    return dict(
        mcq=out,
        short=[dict(unit=str(q["unit"]).upper(), question=q["question"].strip()) for q in paper["short"]],
        long=[dict(unit=str(q["unit"]).upper(), question=q["question"].strip(), marks=str(q["marks"]).strip())
              for q in paper["long"]],
    )


def generate_set(llm, course, units, plan, prior_questions, short_words, long_words, max_tries=4, log=None):
    unit_ids = [u["no"] for u in units]
    errors = None
    for attempt in range(1, max_tries + 1):
        prompt = build_prompt(course, units, plan, prior_questions, short_words, long_words, errors)
        try:
            paper = parse_json(llm(SYSTEM_PROMPT, prompt))
            errors = validate_paper(paper, plan, prior_questions, unit_ids)
        except Exception as e:  # bad JSON, network hiccup, rate limit
            errors = [f"Technical problem: {e}"]
        if log:
            log(f"Set {plan['set_no']} attempt {attempt}: " + ("passed all checks" if not errors else f"{len(errors)} issue(s)"))
        if not errors:
            return finalise(paper), attempt, []
    return None, max_tries, errors


# --------------------------------------------------------------------------
# Step 5: fill the Word template
# --------------------------------------------------------------------------

def _set_cell(cell, text):
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    if p.runs:
        run = p.runs[0]
        for r in p.runs[1:]:
            r._r.getparent().remove(r._r)
        run.text = ""
    else:
        run = p.add_run("")
        ppr_rpr = p._p.find(qn("w:pPr") + "/" + qn("w:rPr"))
        if ppr_rpr is not None:
            old = run._r.find(qn("w:rPr"))
            if old is not None:
                run._r.remove(old)
            run._r.insert(0, copy.deepcopy(ppr_rpr))
        else:
            run.font.name = "Times New Roman"
    for i, line in enumerate(text.split("\n")):
        if i:
            run.add_break()
        run.add_text(line)


def _first_data_row(table):
    for i, r in enumerate(table.rows):
        if r.cells[0].text.strip() == "1":
            return i
    raise ValueError("Template table has no row numbered 1")


def build_docx(template_bytes: bytes, paper: dict, set_no: int, meta: dict) -> bytes:
    d = docx.Document(io.BytesIO(template_bytes))
    if len(d.tables) < 4:
        raise ValueError("Template must contain the header, MCQ, short and long tables")

    for p in d.paragraphs:
        if "Set No" in p.text:
            for r in p.runs:
                if r.text.strip() in ("\u2013", "-", "--", "\u2014"):
                    r.text = str(set_no)

    labels = {"name of programme": meta["programme"], "semester": meta["semester"],
              "course code": meta["code"], "course name": meta["course"],
              "paper setter name": meta.get("setter", "")}
    for row in d.tables[0].rows:
        key = row.cells[0].text.strip().rstrip(":").lower()
        if key in labels and labels[key]:
            _set_cell(row.cells[1], labels[key])

    t = d.tables[1]
    s = _first_data_row(t)
    for i, q in enumerate(paper["mcq"]):
        c = t.rows[s + i].cells
        _set_cell(c[1], q["question"])
        for k in range(4):
            _set_cell(c[2 + k], q["options"][k])
        _set_cell(c[6], str(q["correct"]))
        _set_cell(c[7], q["unit"])

    t = d.tables[2]
    s = _first_data_row(t)
    for i, q in enumerate(paper["short"]):
        c = t.rows[s + i].cells
        _set_cell(c[1], q["question"]); _set_cell(c[2], "5"); _set_cell(c[3], q["unit"])

    t = d.tables[3]
    s = _first_data_row(t)
    for i, q in enumerate(paper["long"]):
        c = t.rows[s + i].cells
        _set_cell(c[1], q["question"]); _set_cell(c[2], q["marks"]); _set_cell(c[3], q["unit"])

    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def zip_files(files: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()
